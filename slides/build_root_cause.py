from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

S = 13.333 / 1131  # inches per source-pixel
def px(v): return Inches(v * S)
def C(h): return RGBColor.from_string(h)

NAVY, BLUE, YEL = "0A4AA6", "0B63F6", "FFC000"
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
sl = prs.slides.add_slide(prs.slide_layouts[6])

def box(x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75, radius=None):
    s = sl.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = C(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    return s

def text(shape, runs, size=10, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=4):
    """runs: list of paragraphs; each paragraph = list of (text, dict(bold, italic, color, size))"""
    tf = shape.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(margin); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, f in para:
            import re as _re
            for seg in _re.split(r"(\[\d\])", t):
                if not seg: continue
                r = p.add_run(); r.text = seg; r.font.name = "Arial"
                if _re.fullmatch(r"\[\d\]", seg): r.font._element.set("baseline", "30000")
                r.font.size = Pt(f.get("size", size)); r.font.bold = f.get("bold", False)
                r.font.italic = f.get("italic", False); r.font.color.rgb = C(f.get("color", "000000"))
            continue
            r = p.add_run(); r.text = t; r.font.name = "Arial"
            r.font.size = Pt(f.get("size", size)); r.font.bold = f.get("bold", False)
            r.font.italic = f.get("italic", False); r.font.color.rgb = C(f.get("color", "000000"))
    return shape

def tb(x, y, w, h, runs, **kw):
    return text(box(x, y, w, h), runs, **kw)

# ---- Title ----
tb(28, 12, 700, 26, [[("Root-Cause Analysis", dict(bold=True, italic=True, color=BLUE, size=18))]], margin=0)
tb(28, 36, 800, 26, [[("Three root causes", dict(bold=True, color=BLUE, size=16)),
                      (" — not awareness or commission rate — explain the adoption gap", dict(size=16))]], margin=0)
tb(840, 14, 130, 40, [[("CAREER", dict(bold=True, size=10))], [("INSIGHT", dict(bold=True, size=10))]], margin=0)
box(975, 16, 1, 36, fill="BBBBBB")
tb(990, 14, 120, 40, [[("blibli", dict(bold=True, color=BLUE, size=22))]], margin=0)

# ---- Key question panel ----
box(28, 104, 164, 423, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
tb(40, 114, 140, 16, [[("KEY QUESTION", dict(bold=True, color=YEL, size=9))]], margin=0)
tb(40, 134, 142, 120, [[("Why does Blibli Affiliate win so little creator effort despite its trust and ecosystem assets?",
                         dict(bold=True, italic=True, color="FFFFFF", size=12))]], anchor=MSO_ANCHOR.TOP, margin=0)
tb(40, 400, 144, 118, [[("The tree is MECE on a creator's expected value per post: ", dict(color="FFFFFF", size=8)),
                        ("buyers reached (A)", dict(bold=True, color=YEL, size=8)), (" × ", dict(color="FFFFFF", size=8)),
                        ("pay per buyer (B)", dict(bold=True, color=YEL, size=8)), (" × ", dict(color="FFFFFF", size=8)),
                        ("confidence to act (C)", dict(bold=True, color=YEL, size=8)), ("[1]", dict(color="FFFFFF", size=6))]],
   anchor=MSO_ANCHOR.BOTTOM, margin=0)

# ---- Branches + leaves ----
groups = [
 ("A", "Demand access", "creators can't turn their audience into Blibli buyers", 113, [
   ("Few followers buy on Blibli: ≈1–3% of physical-goods e-commerce GMV[2]", "E* U"),
   ("Its strongest categories are bought rarely: ≈0.65 orders per user a month[2]", "E*"),
   ("TikTok keeps commerce in-app; Blibli's own affiliate channels are YouTube, IG, WhatsApp", "U P")]),
 ("B", "Earnings design", "the upside is Blibli-funded, capped and narrow", 252, [
   ("No scalable partner-funded upside: Shopee's seller XTRA is uncapped; Blibli's brand bonus is small[3]", "P"),
   ("Narrow earnable surface: ≈39–52% of online B2C TPV earns nothing; cross-vertical buys void[3]", "E* P"),
   ("Short credit on travel (24 h) and two separate balances[3]", "P")]),
 ("C", "Rules & trust", "grassroots creators can't predict or defend earnings", 392, [
   ("Circle selling is voidable: \"relasi\", resale, dropship and bulk purchases excluded[3]", "P"),
   ("\"Affiliate\" carries scam stigma — Blibli repeatedly warns of fake schemes[5]", "P U"),
   ("No public level, sample or progression ladder", "P M")]),
]
spine = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(200), px(176), px(200), px(455))
spine.line.color.rgb = C(NAVY); spine.line.width = Pt(1.25)
for letter, title, sub, y, leaves in groups:
    mid = y + 63
    for (x1, y1, x2, y2) in [(192, mid, 200, mid)] if False else []: pass
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(192), px(mid), px(211), px(mid))
    c.line.color.rgb = C(NAVY); c.line.width = Pt(1.25)
    b = box(211, y, 169, 126, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    circ = box(221, y + 11, 24, 24, fill=YEL, shape=MSO_SHAPE.OVAL)
    text(circ, [[(letter, dict(bold=True, color=NAVY, size=11))]], align=PP_ALIGN.CENTER, margin=0)
    tb(250, y + 8, 125, 28, [[(title, dict(bold=True, color="FFFFFF", size=12))]], margin=0)
    tb(221, y + 42, 152, 70, [[(sub, dict(color="FFFFFF", size=9))]], anchor=MSO_ANCHOR.TOP, margin=0)
    for i, (t, tag) in enumerate(leaves):
        ly = y + i * 44
        l = box(398, ly, 460, 38, fill="FFFFFF", line="C9D3E3", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        tb(406, ly, 385, 38, [[(t, dict(size=8.5))]], margin=0)
        pill = box(805, ly + 8, 45, 22, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        text(pill, [[(tag, dict(bold=True, color="FFFFFF", size=8))]], align=PP_ALIGN.CENTER, margin=0)
        c2 = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(380), px(mid), px(398), px(ly + 19))
        c2.line.color.rgb = C(NAVY); c2.line.width = Pt(0.75)

# ---- Rejected causes ----
tb(872, 104, 230, 20, [[("Rejected as the primary cause", dict(bold=True, color="D93A3A", size=10))]], margin=0)
rej = [("Low awareness", "Moderate, not root: unmeasured, entry is easy — and A's in-app prompt reaches buyers anyway", 126, 52),
       ("Registration friction", "No follower minimum; links work before KYC[3]", 185, 47),
       ("Commission rate too low", "Blibli's caps (Rp30–150k) sit above Shopee's Rp10k base cap[4]", 244, 53)]
for t, d, y, h in rej:
    box(872, y, 230, h, fill="FBE4E4", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    x = box(881, y + 10, 18, 18, fill="D93A3A", shape=MSO_SHAPE.OVAL)
    text(x, [[("✕", dict(bold=True, color="FFFFFF", size=8))]], align=PP_ALIGN.CENTER, margin=0)
    tb(906, y + 4, 192, h - 6, [[(t, dict(bold=True, size=9))], [(d, dict(size=8))]], anchor=MSO_ANCHOR.TOP, margin=0)

# ---- Five Whys ----
box(872, 311, 230, 215, fill="FFF4C7", line=YEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
tb(884, 320, 210, 200, [[("Five Whys — all three chains converge[6]", dict(bold=True, color=NAVY, size=9.5))],
    [("", dict(size=4))],
    [("Affiliate is run as a Blibli-funded performance channel on a mostly first-party, travel-heavy ecosystem, "
      "with rules written for fraud control rather than for creators. The assets exist; the operating model doesn't let creators use them.",
      dict(size=9))]], anchor=MSO_ANCHOR.TOP, margin=0)

# ---- Restated banner ----
b = box(28, 537, 1074, 56, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
tb(40, 537, 1050, 56, [[("Restated: ", dict(bold=True, color=YEL, size=10)),
    ("the gap is driven by (A) demand access — audiences rarely buy on Blibli and its best categories are bought rarely; "
     "(B) earnings design — the upside is Blibli-funded, capped and narrow; (C) rules & trust — circle selling is voidable and \"affiliate\" carries scam stigma.",
     dict(color="FFFFFF", size=10))]], margin=0)

# ---- Footer ----
tb(28, 600, 700, 24, [[("Source: [1] Issue tree, Appendix 18  [2] GDN 1H26 results; Appendix 10  [3] Blibli Tiket Affiliate S&K (2026); Appendix 20  "
                        "[4] Shopee Help Centre (2025)  [5] Blibli advisories; Appendix 21  [6] Five Whys & qualification, Appendix 19",
                        dict(italic=True, color="666666", size=6.5))]], anchor=MSO_ANCHOR.TOP, margin=0)
nav = [("Background", False), ("Analysis", True), ("Solution", False), ("Implementation", False), ("Projection", False)]
tb(740, 606, 330, 14, [[(n + ("    " if i < 4 else ""), dict(size=7, bold=a, color=BLUE if a else "AAAAAA"))
                        for i, (n, a) in enumerate(nav)]], margin=0)
tb(1075, 600, 27, 20, [[("7", dict(bold=True, size=11))]], align=PP_ALIGN.RIGHT, margin=0)

for sh in sl.shapes:
    st = sh._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: sh._element.remove(st)
prs.save("/home/user/ai-consulting-assistant/slides/root_cause_analysis.pptx")
